"""
Build the Word version of KPD_Funding_Playbook_and_Master_Proposal.md
(cover page, table of contents, page numbers). Re-run after editing the .md file.

Usage:
    python build_master_document.py
"""

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from make_word_templates import add_markdown, set_fonts

HERE = Path(__file__).parent
SRC = HERE / "KPD_Funding_Playbook_and_Master_Proposal.md"
OUT = HERE / "KPD_Funding_Playbook_and_Master_Proposal.docx"
NAVY = RGBColor(0x12, 0x3C, 0x69)


def add_field(run, instr):
    for tag, text in (("begin", None), (None, instr), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def cover(doc):
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run("Kon Pop Detection")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(32), NAVY
    p = doc.add_paragraph()
    r = p.add_run("Funding Playbook & Master Proposal")
    r.font.size, r.font.color.rgb = Pt(22), NAVY
    p = doc.add_paragraph()
    r = p.add_run("AI-assisted early oral cancer screening by village health volunteers, Health Region 10, Thailand")
    r.italic, r.font.size = True, Pt(13)
    doc.add_paragraph()
    for line in (
        f"Prepared: {date.today():%d %B %Y}",
        "For: the Kon Pop Detection team, partner hospitals and Mahidol University collaborators",
        "Status: working draft. Replace every [FILL]; verify every [CONFIRM] against the official call.",
    ):
        doc.add_paragraph(line)
    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run("Contents of this document: funding strategy · calendar · funder directory · what proposals must contain · "
                  "master proposal · budget guide · evidence · compliance · registrations · letters · fixes · glossary")
    r.font.size, r.font.color.rgb = Pt(9.5), RGBColor(0x55, 0x55, 0x55)
    doc.add_page_break()


def main():
    doc = Document()
    set_fonts(doc)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(2.0)
        fp = s.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.add_run("Kon Pop Detection · Funding Playbook · page ").font.size = Pt(8)
        add_field(fp.add_run(), "PAGE")
    # ask Word to refresh the table of contents when the file is opened
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    doc.settings.element.append(upd)

    cover(doc)
    add_markdown(doc, SRC.read_text(encoding="utf-8").splitlines())
    doc.save(OUT)
    print(f"Written: {OUT}")


if __name__ == "__main__":
    main()
