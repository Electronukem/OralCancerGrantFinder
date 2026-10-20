"""
Convert the Markdown files in a folder (default ./templates) into fillable Word (.docx)
files in <folder>/word. Re-run after editing any .md file.

Usage:
    python make_word_templates.py              # templates
    python make_word_templates.py applications # application drafts
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

SRC = Path(__file__).parent / (sys.argv[1] if len(sys.argv) > 1 else "templates")
OUT = SRC / "word"
THAI_FONT = "Leelawadee UI"   # renders Thai text correctly
LATIN_FONT = "Calibri"

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def set_fonts(doc):
    style = doc.styles["Normal"]
    style.font.name = LATIN_FONT
    style.font.size = Pt(10.5)
    rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:cs"), THAI_FONT)
    rfonts.set(qn("w:eastAsia"), THAI_FONT)
    for name in ("Heading 1", "Heading 2", "Heading 3", "Title"):
        rf = doc.styles[name].element.get_or_add_rPr().get_or_add_rFonts()
        rf.set(qn("w:cs"), THAI_FONT)


def add_inline(par, text):
    """Add text to a paragraph, honouring **bold** and `code`."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
        else:
            par.add_run(part)


def cells(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


HEADER_FILL = "DCE6F1"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_table(doc, rows):
    header, body = rows[0], [r for r in rows[1:] if not re.fullmatch(r"[\s:\-|]+", "|".join(r))]
    ncols = len(header)
    t = doc.add_table(rows=1, cols=ncols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        shade(cell, HEADER_FILL)
        p = cell.paragraphs[0]
        add_inline(p, h)
        for r in p.runs:
            r.bold = True
    for row in body:
        tc = t.add_row().cells
        for i in range(ncols):
            add_inline(tc[i].paragraphs[0], row[i] if i < len(row) else "")
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9.5)
    doc.add_paragraph()


def add_toc(doc):
    """Insert a Word table-of-contents field (Word fills it in when the document is opened/updated)."""
    p = doc.add_paragraph()
    run = p.add_run()
    for tag, text in (("begin", None), (None, 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
            if tag == "begin":
                el.set(qn("w:dirty"), "true")
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
        if tag == "separate":
            run._r.append(OxmlElement("w:t"))
            run._r[-1].text = "Right-click here and choose 'Update Field' to show the table of contents."


def convert(md_path, out_path):
    doc = Document()
    set_fonts(doc)
    add_markdown(doc, md_path.read_text(encoding="utf-8").splitlines())
    doc.save(out_path)


def add_markdown(doc, lines):
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(cells(lines[i]))
                i += 1
            add_table(doc, rows)
            continue
        if s == "<!-- pagebreak -->":
            doc.add_page_break()
        elif s == "<!-- toc -->":
            doc.add_heading("Contents", level=1)
            add_toc(doc)
        elif not s or s == "---" or s.startswith("<!--"):
            pass
        elif s.startswith("# "):
            doc.add_heading(s[2:], level=1)
        elif s.startswith("## "):
            doc.add_heading(s[3:], level=2)
        elif s.startswith("### "):
            doc.add_heading(s[4:], level=3)
        elif s.startswith("#### "):
            doc.add_heading(s[5:], level=4)
        elif s.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            add_inline(p, s.lstrip("> ").strip())
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = RGBColor(0x33, 0x44, 0x66)
        elif re.match(r"^[-*] ", s):
            indent = (len(line) - len(line.lstrip())) // 2
            p = doc.add_paragraph(style="List Bullet 2" if indent else "List Bullet")
            add_inline(p, s[2:])
        elif re.match(r"^\d+\. ", s):
            # keep the written number (Word's auto-numbering would continue across sections)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            p.paragraph_format.first_line_indent = Pt(-14)
            add_inline(p, s)
        else:
            add_inline(doc.add_paragraph(), s)
        i += 1


def main():
    OUT.mkdir(exist_ok=True)
    files = sorted(SRC.glob("*.md"))
    for f in files:
        out = OUT / (f.stem + ".docx")
        convert(f, out)
        print(f"  {out.name}")
    print(f"\n{len(files)} Word templates written to {OUT}")


if __name__ == "__main__":
    main()
